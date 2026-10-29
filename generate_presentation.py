"""
MutationHunter PowerPoint Presentation Generator v2 (Executive / High-Design Edition)
Creates a modern, award-winning 16:9 presentation deck adhering to top-tier academic
and biotech conference presentation standards.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_deck(output_path="MutationHunter_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # COLOR PALETTE
    # =========================================================================
    NAVY_DARK    = RGBColor(11, 19, 43)     # #0B132B Dark Navy Header / Title
    NAVY_SURFACE = RGBColor(15, 23, 42)     # #0F172A
    CARD_BG      = RGBColor(255, 255, 255)  # Pure White
    PAGE_BG      = RGBColor(248, 250, 252)  # #F8FAFC Soft Canvas Slate
    BORDER_LIGHT = RGBColor(226, 232, 240)  # #E2E8F0 Subtle Border
    BORDER_DARK  = RGBColor(203, 213, 225)  # #CBD5E1

    TEXT_TITLE   = RGBColor(15, 23, 42)     # #0F172A
    TEXT_BODY    = RGBColor(51, 65, 85)     # #334155
    TEXT_MUTED   = RGBColor(100, 116, 139)  # #64748B
    TEXT_LIGHT   = RGBColor(148, 163, 184)  # #94A3B8

    # Accent Hues
    BLUE_PRIMARY = RGBColor(37, 99, 235)    # #2563EB Royal Blue
    BLUE_LIGHT   = RGBColor(239, 246, 255)  # #EFF6FF Tint
    BLUE_DARK    = RGBColor(29, 78, 216)    # #1D4ED8

    EMERALD      = RGBColor(5, 150, 105)    # #059669 Emerald Green
    EMERALD_LT   = RGBColor(236, 253, 245)  # #ECFDF5 Tint
    
    INDIGO       = RGBColor(79, 70, 229)    # #4F46E5 Indigo
    INDIGO_LT    = RGBColor(238, 242, 255)  # #EEF2FF Tint

    AMBER        = RGBColor(217, 119, 6)    # #D97706 Amber
    AMBER_LT     = RGBColor(254, 243, 199)  # #FEF3C7 Tint

    CORAL        = RGBColor(220, 38, 38)    # #DC2626 Coral Red
    CORAL_LT     = RGBColor(254, 242, 242)  # #FEF2F2 Tint

    FONT_MAIN = "Segoe UI"

    # =========================================================================
    # REUSABLE UI HELPERS
    # =========================================================================
    def apply_canvas(slide, dark=False):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY_SURFACE if dark else PAGE_BG
        bg.line.fill.background()
        return bg

    def add_top_nav(slide, category_badge, title_text, subtitle_text, slide_num):
        # 1. Subtle top primary accent line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.06))
        line.fill.solid()
        line.fill.fore_color.rgb = BLUE_PRIMARY
        line.line.fill.background()

        # 2. Category Pill Badge
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(3.2), Inches(0.28))
        pill.fill.solid()
        pill.fill.fore_color.rgb = BLUE_LIGHT
        pill.line.color.rgb = RGBColor(191, 219, 254)
        pill.line.width = Pt(1)
        p = pill.text_frame.paragraphs[0]
        p.text = category_badge.upper()
        p.font.name = FONT_MAIN
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = BLUE_PRIMARY
        p.alignment = PP_ALIGN.CENTER

        # 3. Slide Number Pill (Top Right)
        num_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.333), Inches(0.35), Inches(1.2), Inches(0.28))
        num_pill.fill.solid()
        num_pill.fill.fore_color.rgb = RGBColor(241, 245, 249)
        num_pill.line.color.rgb = BORDER_LIGHT
        num_pill.line.width = Pt(1)
        pn = num_pill.text_frame.paragraphs[0]
        pn.text = f"{slide_num:02d} / 10"
        pn.font.name = FONT_MAIN
        pn.font.size = Pt(9)
        pn.font.bold = True
        pn.font.color.rgb = TEXT_MUTED
        pn.alignment = PP_ALIGN.CENTER

        # 4. Slide Title & Subtitle
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(10.5), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title_text
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_TITLE

        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(2)

        # 5. Bottom subtle footer
        add_footer(slide)

    def add_footer(slide):
        # Footer dividing line
        f_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.015))
        f_line.fill.solid()
        f_line.fill.fore_color.rgb = BORDER_LIGHT
        f_line.line.fill.background()

        # Left label
        tb_f = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(8.0), Inches(0.3))
        tf_f = tb_f.text_frame
        tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
        p = tf_f.paragraphs[0]
        p.text = "MutationHunter | Nature-Inspired Computing (GA & PSO) in Clinical Genomics"
        p.font.name = FONT_MAIN
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_MUTED

        # Right label
        tb_fr = slide.shapes.add_textbox(Inches(9.533), Inches(7.1), Inches(3.0), Inches(0.3))
        tf_fr = tb_fr.text_frame
        tf_fr.margin_left = tf_fr.margin_top = tf_fr.margin_right = tf_fr.margin_bottom = 0
        pr = tf_fr.paragraphs[0]
        pr.text = "ClinVar Conflicting Variant Study"
        pr.font.name = FONT_MAIN
        pr.font.size = Pt(9)
        pr.font.color.rgb = TEXT_MUTED
        pr.alignment = PP_ALIGN.RIGHT

    def create_card(slide, left, top, width, height, accent_color=None, fill_color=CARD_BG, border_color=BORDER_LIGHT):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        # Optional colored top accent stripe
        if accent_color:
            stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.02), top + Inches(0.02), width - Inches(0.04), Inches(0.06))
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = accent_color
            stripe.line.fill.background()
        return card

    def add_kpi_card(slide, left, top, width, height, value_text, label_text, sub_text="", val_color=BLUE_PRIMARY, bg_color=CARD_BG, border_color=BORDER_LIGHT):
        create_card(slide, left, top, width, height, accent_color=val_color, fill_color=bg_color, border_color=border_color)
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), height - Inches(0.25))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = value_text
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = val_color

        p2 = tf.add_paragraph()
        p2.text = label_text
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(10.5)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_TITLE
        p2.space_before = Pt(2)

        if sub_text:
            p3 = tf.add_paragraph()
            p3.text = sub_text
            p3.font.name = FONT_MAIN
            p3.font.size = Pt(9)
            p3.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 1: Title Slide (Executive Deep Navy Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_canvas(s1, dark=True)

    # Decorative geometric background elements
    deco1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
    deco1.fill.solid()
    deco1.fill.fore_color.rgb = BLUE_PRIMARY
    deco1.line.fill.background()

    # Hero Card Container
    hero = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.8), Inches(11.533), Inches(5.9))
    hero.fill.solid()
    hero.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
    hero.line.color.rgb = RGBColor(30, 41, 59)
    hero.line.width = Pt(1.5)

    # Category Pill
    pill1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(1.3), Inches(4.5), Inches(0.35))
    pill1.fill.solid()
    pill1.fill.fore_color.rgb = RGBColor(30, 58, 138) # Dark blue tint
    pill1.line.color.rgb = RGBColor(59, 130, 246)
    p = pill1.text_frame.paragraphs[0]
    p.text = "RESEARCH CAPSTONE | NATURE-INSPIRED COMPUTING"
    p.font.name = FONT_MAIN
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    p.alignment = PP_ALIGN.CENTER

    # Project Title
    tb_title = s1.shapes.add_textbox(Inches(1.4), Inches(1.8), Inches(10.5), Inches(1.6))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "MutationHunter: Nature-Inspired Detection & Prioritization of Conflicting Genetic Variant Classifications"
    p.font.name = FONT_MAIN
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = tf_t.add_paragraph()
    p_sub.text = "Combinatorial Biomarker Optimization & Discordance Classification Using Genetic Algorithms (GA) and Particle Swarm Optimization (PSO)"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_sub.space_before = Pt(8)

    # Sub-Grid: Team Card & Mentor Card
    # Team Card
    card_team = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(3.8), Inches(5.1), Inches(2.4))
    card_team.fill.solid()
    card_team.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card_team.line.color.rgb = RGBColor(51, 65, 85)
    
    tb_team = s1.shapes.add_textbox(Inches(1.65), Inches(3.95), Inches(4.6), Inches(2.1))
    tf_team = tb_team.text_frame
    tf_team.word_wrap = True
    p = tf_team.paragraphs[0]
    p.text = "PROJECT TEAM MEMBERS"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248) # Sky blue

    team_members = [
        ("Student 1", "[Student Name 1]", "Roll No: [Roll Number 1]"),
        ("Student 2", "[Student Name 2]", "Roll No: [Roll Number 2]"),
        ("Student 3", "[Student Name 3]", "Roll No: [Roll Number 3]")
    ]
    for role, name, roll in team_members:
        p_mem = tf_team.add_paragraph()
        p_mem.text = f"•  {name}  —  {roll}"
        p_mem.font.name = FONT_MAIN
        p_mem.font.size = Pt(11)
        p_mem.font.color.rgb = RGBColor(241, 245, 249)
        p_mem.space_before = Pt(4)

    # Guide Card
    card_guide = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.8), Inches(5.1), Inches(2.4))
    card_guide.fill.solid()
    card_guide.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card_guide.line.color.rgb = RGBColor(51, 65, 85)

    tb_guide = s1.shapes.add_textbox(Inches(7.05), Inches(3.95), Inches(4.6), Inches(2.1))
    tf_guide = tb_guide.text_frame
    tf_guide.word_wrap = True
    p = tf_guide.paragraphs[0]
    p.text = "ACADEMIC MENTORSHIP & INSTITUTION"
    p.font.name = FONT_MAIN
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(52, 211, 153) # Mint green

    guide_items = [
        "Project Guide: [Guide Name / Faculty Advisor]",
        "Designation: [Associate Professor / Professor]",
        "Department: Department of Computer Science & Engineering",
        "Institution: [College / University Name]"
    ]
    for g in guide_items:
        p_g = tf_guide.add_paragraph()
        p_g.text = f"•  {g}"
        p_g.font.name = FONT_MAIN
        p_g.font.size = Pt(11)
        p_g.font.color.rgb = RGBColor(241, 245, 249)
        p_g.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Problem Statement & Clinical Motivation
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_canvas(s2)
    add_top_nav(s2, "Context & Clinical Bottleneck", "Slide 2: Problem Statement & Clinical Motivation", "Addressing diagnostic discordance in multi-submitter genetic variant repositories", 2)

    # 3 High-Impact KPI badges
    add_kpi_card(s2, Inches(0.8), Inches(1.5), Inches(3.7), Inches(1.1), "25.2%", "Clinical Discordance Rate", "1 in 4 variants receive conflicting classifications", val_color=CORAL, bg_color=CORAL_LT)
    add_kpi_card(s2, Inches(4.8), Inches(1.5), Inches(3.7), Inches(1.1), "46+ Features", "Genomic Dimensionality", "Collinear in silico tools obscure true drivers", val_color=AMBER, bg_color=AMBER_LT)
    add_kpi_card(s2, Inches(8.8), Inches(1.5), Inches(3.7), Inches(1.1), "Months Delay", "Manual Consensus Review", "Expert panels require extensive committee time", val_color=BLUE_PRIMARY, bg_color=BLUE_LIGHT)

    # 3 Column Detail Cards
    col_y = Inches(2.8)
    col_w = Inches(3.7)
    col_h = Inches(4.0)

    # Card 1: What problem are you solving?
    create_card(s2, Inches(0.8), col_y, col_w, col_h, accent_color=CORAL)
    tb1 = s2.shapes.add_textbox(Inches(1.0), col_y + Inches(0.2), col_w - Inches(0.4), col_h - Inches(0.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "What Problem Are We Solving?"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    p_sub = tf1.add_paragraph()
    p_sub.text = "Discordant Variant Interpretations"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = CORAL
    p_sub.space_before = Pt(2)

    bullets1 = [
        "Independent diagnostic laboratories (Ambry, GeneDx, Invitae) submit contradictory clinical calls for identical variants.",
        "A single mutation may be labelled 'Pathogenic' by one lab and 'Benign' or 'VUS' by another.",
        "High-dimensional tables have heavy multicollinearity across CADD, SIFT, and PolyPhen predictors."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(6)

    # Card 2: Why is it important?
    create_card(s2, Inches(4.8), col_y, col_w, col_h, accent_color=AMBER)
    tb2 = s2.shapes.add_textbox(Inches(5.0), col_y + Inches(0.2), col_w - Inches(0.4), col_h - Inches(0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Why Is It Important?"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    p_sub = tf2.add_paragraph()
    p_sub.text = "Severe Healthcare Consequences"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = AMBER
    p_sub.space_before = Pt(2)

    bullets2 = [
        "In oncology, discordant calls lead to either missing lethal malignancies or administering toxic unnecessary chemotherapies.",
        "In medical genetics, patients undergo extended diagnostic odysseys without clear therapeutic pathways.",
        "Manual consensus requires months of multi-institutional expert panel review per contested variant."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(6)

    # Card 3: Who will benefit?
    create_card(s2, Inches(8.8), col_y, col_w, col_h, accent_color=EMERALD)
    tb3 = s2.shapes.add_textbox(Inches(9.0), col_y + Inches(0.2), col_w - Inches(0.4), col_h - Inches(0.4))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Who Will Benefit?"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    p_sub = tf3.add_paragraph()
    p_sub.text = "Genomic Medicine Stakeholders"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(10)
    p_sub.font.bold = True
    p_sub.font.color.rgb = EMERALD
    p_sub.space_before = Pt(2)

    bullets3 = [
        "Clinical Testing Labs: Automated quality-control filter flagging high-conflict mutations before sign-out.",
        "Genetic Counselors: Actionable biological explainability showing exactly why classifications diverge.",
        "Patients: Faster turnaround, accurate treatment selection, and greater diagnostic certainty."
    ]
    for b in bullets3:
        p = tf3.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 3: Objectives & Success Criteria
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_canvas(s3)
    add_top_nav(s3, "Goals & Evaluation Criteria", "Slide 3: Project Objectives & Expected Outcomes", "Navigating 1.07 billion combinatorial subsets to evolve interpretable diagnostic panels", 3)

    quad_w = Inches(5.7)
    quad_h = Inches(2.55)

    # Quadrant 1: Main Project Goal
    create_card(s3, Inches(0.8), Inches(1.6), quad_w, quad_h, accent_color=BLUE_PRIMARY)
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(1.8), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Main Project Goal"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    bullets = [
        "Formulate Nature-Inspired Computing (NIC) algorithms to navigate the 2^30 (~1.07 Billion) combinatorial feature search space in ClinVar.",
        "Prune redundant and collinear genomic attributes while retaining high-fidelity biophysical biomarkers.",
        "Detect and prioritize conflicting clinical variant assertions (CLASS=1) versus consistent submissions (CLASS=0)."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Quadrant 2: Proposed Solution
    create_card(s3, Inches(6.8), Inches(1.6), quad_w, quad_h, accent_color=INDIGO)
    tb = s3.shapes.add_textbox(Inches(7.0), Inches(1.8), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Proposed Solution"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = INDIGO

    bullets = [
        "Implement a Darwinian Genetic Algorithm (GA) with tournament selection, uniform crossover, and bit-flip mutation.",
        "Implement Binary Particle Swarm Optimization (BPSO) with velocity clamping and sigmoid transfer mapping.",
        "Train cost-sensitive balanced Random Forest ensembles on the evolved biomarker panels to overcome 75/25 class imbalance."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Quadrant 3: Expected Outcome
    create_card(s3, Inches(0.8), Inches(4.35), quad_w, quad_h, accent_color=EMERALD)
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(4.55), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Expected Outcome"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    bullets = [
        "Achieve over 70% feature dimensionality reduction while matching or exceeding full 30-feature baseline accuracy.",
        "Uncover biologically verified drivers of discordance (gene loss-of-function intolerance, maximum allele frequency, high consequence impact).",
        "Significantly accelerate inference latency to enable real-time clinical integration."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Quadrant 4: Quantitative Success Criteria
    create_card(s3, Inches(6.8), Inches(4.35), quad_w, quad_h, accent_color=AMBER)
    tb = s3.shapes.add_textbox(Inches(7.0), Inches(4.55), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "4. Quantitative Success Criteria"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AMBER

    bullets = [
        "Classification Accuracy >= 70.0% and ROC-AUC >= 0.76 on held-out test split of 13,038 variants.",
        "Feature Sparsity: >= 80% reduction with GA and >= 70% reduction with BPSO.",
        "Inference Speed: Sub-50 millisecond per-variant classification latency.",
        "Interactive Deployment: Functional clinical decision dashboard with live explainability."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 4: Dataset & System Architecture
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_canvas(s4)
    add_top_nav(s4, "Data & Architecture", "Slide 4: Dataset Specifications & System Architecture", "NCBI ClinVar 65,188 human genomic variant repository and end-to-end pipeline", 4)

    # Left: Dataset Profile Card & Mini-KPIs
    left_w = Inches(5.3)
    create_card(s4, Inches(0.8), Inches(1.5), left_w, Inches(5.3), accent_color=BLUE_PRIMARY)
    
    tb = s4.shapes.add_textbox(Inches(1.0), Inches(1.7), left_w - Inches(0.4), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "ClinVar Conflicting Dataset Profile"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    # 4 Mini-KPIs inside left card
    add_kpi_card(s4, Inches(1.0), Inches(2.2), Inches(2.35), Inches(0.95), "65,188", "Human Variants", "Verified ClinVar cohort", val_color=BLUE_PRIMARY, bg_color=PAGE_BG)
    add_kpi_card(s4, Inches(3.55), Inches(2.2), Inches(2.35), Inches(0.95), "30 Features", "Engineered Inputs", "Across 5 clinical domains", val_color=INDIGO, bg_color=PAGE_BG)
    add_kpi_card(s4, Inches(1.0), Inches(3.3), Inches(2.35), Inches(0.95), "74.8 / 25.2%", "Class Imbalance", "48.7k Consist. / 16.4k Conf.", val_color=AMBER, bg_color=PAGE_BG)
    add_kpi_card(s4, Inches(3.55), Inches(3.3), Inches(2.35), Inches(0.95), "80 / 20 Split", "Stratified Partition", "52,150 Train / 13,038 Test", val_color=EMERALD, bg_color=PAGE_BG)

    tb_ds = s4.shapes.add_textbox(Inches(1.0), Inches(4.4), left_w - Inches(0.4), Inches(2.2))
    tf_ds = tb_ds.text_frame
    tf_ds.word_wrap = True
    tf_ds.paragraphs[0].text = "Source & Genomic Annotation Integration:"
    tf_ds.paragraphs[0].font.name = FONT_MAIN
    tf_ds.paragraphs[0].font.size = Pt(10.5)
    tf_ds.paragraphs[0].font.bold = True
    tf_ds.paragraphs[0].font.color.rgb = TEXT_TITLE

    sources = [
        "Dataset Source: NCBI ClinVar (kevinarvai/clinvar-conflicting), MUTATIONHUNTER Rank 1.",
        "Integrated Databases: ExAC, 1000 Genomes (TGP), ESP6500, CADD v1.6, LoFtool intolerance.",
        "Technology Stack: Python 3.14, Scikit-Learn, Pandas, NumPy, Flask, Chart.js, Bootstrap 5."
    ]
    for s in sources:
        p = tf_ds.add_paragraph()
        p.text = f"•  {s}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(3)

    # Right: Visual Pipeline Steps
    right_w = Inches(6.1)
    right_x = Inches(6.4)
    create_card(s4, right_x, Inches(1.5), right_w, Inches(5.3), accent_color=INDIGO)

    tb_rt = s4.shapes.add_textbox(right_x + Inches(0.2), Inches(1.7), right_w - Inches(0.4), Inches(0.4))
    p = tb_rt.text_frame.paragraphs[0]
    p.text = "End-to-End System Pipeline Architecture"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    steps = [
        ("01", "Raw Ingestion", "Stream 65,188 ClinVar variant records with 46 multi-attribute columns.", BLUE_PRIMARY),
        ("02", "Preprocessing & Imputation", "Zero-fill novel population AFs, median-impute CADD/LoFtool, encode consequences.", INDIGO),
        ("03", "Evolutionary Search Cohort", "Extract 5,000 stratified variant sample to enable ultra-rapid fitness evaluation.", EMERALD),
        ("04", "Nature-Inspired Optimization", "Run Darwinian GA (20 gen) and Binary PSO (15 iter) to prune uninformative noise.", AMBER),
        ("05", "Ensemble Modeling & Benchmarking", "Train cost-sensitive balanced Random Forest models on evolved biomarker panels.", CORAL),
        ("06", "Clinical Decision Dashboard", "Interactive Flask web prototype with real-time inference & biological explainability.", BLUE_PRIMARY)
    ]

    step_y = Inches(2.25)
    for num, title, desc, clr in steps:
        # Step Container
        s_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x + Inches(0.2), step_y, right_w - Inches(0.4), Inches(0.68))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = PAGE_BG
        s_box.line.color.rgb = BORDER_LIGHT
        s_box.line.width = Pt(1)

        # Number badge
        nb = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x + Inches(0.3), step_y + Inches(0.12), Inches(0.45), Inches(0.44))
        nb.fill.solid()
        nb.fill.fore_color.rgb = clr
        nb.line.fill.background()
        p = nb.text_frame.paragraphs[0]
        p.text = num
        p.font.name = FONT_MAIN
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        # Step Text
        stb = s4.shapes.add_textbox(right_x + Inches(0.85), step_y + Inches(0.06), right_w - Inches(1.1), Inches(0.55))
        stf = stb.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
        p = stf.paragraphs[0]
        p.text = title
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_TITLE

        p2 = stf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

        step_y += Inches(0.74)

    # =========================================================================
    # SLIDE 5: Nature-Inspired Computing (NIC) Algorithms
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_canvas(s5)
    add_top_nav(s5, "Metaheuristic Optimization", "Slide 5: Nature-Inspired Computing (NIC) Formulations", "Darwinian Genetic Algorithm vs. Binary Particle Swarm Optimization", 5)

    side_w = Inches(5.7)

    # Left: Genetic Algorithm (GA)
    create_card(s5, Inches(0.8), Inches(1.5), side_w, Inches(5.3), accent_color=EMERALD)
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(1.7), side_w - Inches(0.4), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "Genetic Algorithm (GA) - Darwinian Selection"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    # Formula Box
    f_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.2), side_w - Inches(0.4), Inches(0.85))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = EMERALD_LT
    f_box.line.color.rgb = RGBColor(167, 243, 208)
    f_box.line.width = Pt(1)
    tf = f_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Multi-Objective Fitness Formulation:"
    p.font.name = FONT_MAIN
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    p2 = tf.add_paragraph()
    p2.text = "Fitness(c) = 0.92 · AUC_CV(c) + 0.08 · (1 - |c| / 30)"
    p2.font.name = "Consolas"
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_TITLE

    tb_ga = s5.shapes.add_textbox(Inches(1.0), Inches(3.2), side_w - Inches(0.4), Inches(3.4))
    tf_ga = tb_ga.text_frame
    tf_ga.word_wrap = True

    ga_bullets = [
        ("Biological Basis", "Natural selection, survival of the fittest, and genetic mutation."),
        ("Chromosome Encoding", "Binary string c in {0, 1}^30. Bit i = 1 includes feature; 0 prunes it."),
        ("Selection Mechanism", "Tournament Selection (k = 3) favoring high-affinity chromosomes."),
        ("Crossover Operator", "Uniform Crossover (Pc = 0.85) facilitating genetic trait exploration."),
        ("Mutation Operator", "Bit-Flip Mutation (Pm = 0.04) preventing premature convergence."),
        ("Elitism Strategy", "Top 2 elite chromosomes preserved unchanged across 20 generations."),
        ("Evolved Biomarkers", "LoFtool, AF_TGP, Max_AF, IMPACT_HIGH (4 features -> 86.7% reduction!).")
    ]
    for idx, (label, desc) in enumerate(ga_bullets):
        p = tf_ga.paragraphs[0] if idx == 0 else tf_ga.add_paragraph()
        p.text = f"•  {label}:  {desc}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Right: Binary PSO
    create_card(s5, Inches(6.8), Inches(1.5), side_w, Inches(5.3), accent_color=BLUE_PRIMARY)
    tb = s5.shapes.add_textbox(Inches(7.0), Inches(1.7), side_w - Inches(0.4), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "Binary Particle Swarm Optimization (BPSO)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    # Formula Box
    f_box2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(2.2), side_w - Inches(0.4), Inches(0.85))
    f_box2.fill.solid()
    f_box2.fill.fore_color.rgb = BLUE_LIGHT
    f_box2.line.color.rgb = RGBColor(191, 219, 254)
    f_box2.line.width = Pt(1)
    tf2 = f_box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Velocity Update & Sigmoid Transfer Function:"
    p.font.name = FONT_MAIN
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    p2 = tf2.add_paragraph()
    p2.text = "v_i(t+1) = w·v_i + c1·r1·(pbest - x_i) + c2·r2·(gbest - x_i) ;  S(v) = 1/(1+e^-v)"
    p2.font.name = "Consolas"
    p2.font.size = Pt(9.5)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_TITLE

    tb_pso = s5.shapes.add_textbox(Inches(7.0), Inches(3.2), side_w - Inches(0.4), Inches(3.4))
    tf_pso = tb_pso.text_frame
    tf_pso.word_wrap = True

    pso_bullets = [
        ("Biological Basis", "Collective swarm intelligence observed in bird flocking and fish schooling."),
        ("Swarm Topology", "20 particles navigating a 30-dimensional discrete Boolean hypercube."),
        ("Inertia & Social Pull", "Inertia w = 0.729; Cognitive c1 = 1.494; Social c2 = 1.494."),
        ("Velocity Clamping", "Clamped to [-4.0, +4.0] to prevent sigmoid saturation and bit lock."),
        ("Bit Transfer Rule", "Bit x_ij = 1 if rand() < S(v_ij) else 0."),
        ("Convergence Speed", "Converged rapidly in 15 iterations (34.25s runtime)."),
        ("Evolved Biomarkers", "9 Features (LoFtool, AF_EXAC, AF_TGP, Max_AF, CHROM_num, IMPACTs...)."),
        ("Champion Result", "Achieved highest overall accuracy (70.40%) & Macro F1 (0.6602)!")
    ]
    for idx, (label, desc) in enumerate(pso_bullets):
        p = tf_pso.paragraphs[0] if idx == 0 else tf_pso.add_paragraph()
        p.text = f"•  {label}:  {desc}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 6: Research Methodology & Experimental Design
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_canvas(s6)
    add_top_nav(s6, "Methodology & Experiments", "Slide 6: Research Methodology & Experimental Design", "Rigorous 4-stage scientific framework from raw ClinVar curation to clinical validation", 6)

    m_w = Inches(2.75)
    m_h = Inches(5.3)
    m_x_list = [Inches(0.8), Inches(3.78), Inches(6.76), Inches(9.74)]

    stages = [
        ("Phase 1", "Preprocessing & Engineering", BLUE_PRIMARY, [
            "Missing Data Imputation: Zero-fill novel AFs (ESP, ExAC, TGP); median impute CADD and LoFtool.",
            "Composite AF Features: Max_AF = max(AF_ESP, AF_EXAC, AF_TGP) capturing rare vs common bounds.",
            "Novel Variant Flag: Is_Novel = 1 when unobserved across all global population reference cohorts.",
            "Categorical Encoding: One-hot encoded IMPACT, CLNVC, and CSQ consequence categories."
        ]),
        ("Phase 2", "NIC Algorithm Search", INDIGO, [
            "Stratified Search Cohort: 5,000 balanced variants dedicated to rapid metaheuristic evaluation.",
            "Fast Fitness Evaluation: 3-Fold Stratified Cross-Validation on each candidate feature chromosome.",
            "Class Imbalance Mitigation: Balanced class weights preventing majority class dominance.",
            "Convergence Tracking: Real-time generational recording of fitness score and chromosome sparsity."
        ]),
        ("Phase 3", "Model Development", EMERALD, [
            "Full Baseline 1: Random Forest (80 trees, depth 10) trained on all 30 engineered features.",
            "Linear Baseline 2: L2-penalized Logistic Regression on all 30 features.",
            "NIC GA Model: Random Forest trained solely on the 4 GA-evolved biomarkers.",
            "NIC PSO Model: Random Forest trained solely on the 9 PSO-evolved biomarkers."
        ]),
        ("Phase 4", "Validation & Profiling", AMBER, [
            "Strict Held-Out Split: 13,038 unseen ClinVar variants completely isolated from optimization.",
            "Evaluation Metrics: Test Accuracy, Macro F1, Precision, Recall, and ROC-AUC.",
            "Computational Profiling: Search runtime, model training latency, and per-variant inference speed.",
            "Visual Diagnostics: Side-by-side confusion matrices and ROC curve generation."
        ])
    ]

    for idx, (p_num, p_title, p_color, p_bullets) in enumerate(stages):
        x = m_x_list[idx]
        create_card(s6, x, Inches(1.5), m_w, m_h, accent_color=p_color)

        # Header Badge
        h_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), Inches(1.7), Inches(1.2), Inches(0.3))
        h_box.fill.solid()
        h_box.fill.fore_color.rgb = p_color
        h_box.line.fill.background()
        p = h_box.text_frame.paragraphs[0]
        p.text = p_num.upper()
        p.font.name = FONT_MAIN
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        tb = s6.shapes.add_textbox(x + Inches(0.2), Inches(2.1), m_w - Inches(0.4), Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = p_title
        p.font.name = FONT_MAIN
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_TITLE

        tb_b = s6.shapes.add_textbox(x + Inches(0.2), Inches(2.75), m_w - Inches(0.4), Inches(3.9))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        for b_idx, b in enumerate(p_bullets):
            p = tf_b.paragraphs[0] if b_idx == 0 else tf_b.add_paragraph()
            p.text = f"•  {b}"
            p.font.name = FONT_MAIN
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_BODY
            p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 7: Performance Metrics & Quantitative Benchmarks
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_canvas(s7)
    add_top_nav(s7, "Quantitative Benchmarks", "Slide 7: Performance Metrics & Evaluation Profiling", "Evaluation across 13,038 strictly held-out ClinVar test variants", 7)

    # 4 Top Executive Stat Cards
    add_kpi_card(s7, Inches(0.8), Inches(1.5), Inches(2.75), Inches(1.1), "70.40%", "Accuracy (PSO Champion)", "Outperformed 30-feature baseline", val_color=BLUE_PRIMARY, bg_color=BLUE_LIGHT)
    add_kpi_card(s7, Inches(3.78), Inches(1.5), Inches(2.75), Inches(1.1), "86.7%", "Feature Reduction (GA)", "Pruned 26 noisy features down to 4", val_color=EMERALD, bg_color=EMERALD_LT)
    add_kpi_card(s7, Inches(6.76), Inches(1.5), Inches(2.75), Inches(1.1), "0.6602", "Macro F1-Score (PSO)", "Superior detection of minority conflict", val_color=INDIGO, bg_color=INDIGO_LT)
    add_kpi_card(s7, Inches(9.74), Inches(1.5), Inches(2.75), Inches(1.1), "37.7 ms", "Inference Latency", "Ultra-fast clinical decision response", val_color=AMBER, bg_color=AMBER_LT)

    # Benchmark Table Card (Left, 65% width)
    tbl_w = Inches(7.6)
    create_card(s7, Inches(0.8), Inches(2.8), tbl_w, Inches(4.0), accent_color=BLUE_PRIMARY)

    tb = s7.shapes.add_textbox(Inches(1.0), Inches(2.95), tbl_w - Inches(0.4), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = "Comprehensive Benchmark on 13,038 Held-Out Test Variants"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    # Add Table
    tbl_shape = s7.shapes.add_table(5, 7, Inches(1.0), Inches(3.4), tbl_w - Inches(0.4), Inches(2.2))
    table = tbl_shape.table

    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(0.7)
    table.columns[2].width = Inches(0.8)
    table.columns[3].width = Inches(0.8)
    table.columns[4].width = Inches(0.8)
    table.columns[5].width = Inches(0.8)
    table.columns[6].width = Inches(1.2)

    headers = ["Model Name", "Feats", "Accuracy", "Macro F1", "Recall", "ROC-AUC", "Status / Notes"]
    rows = [
        ["Random Forest (Base)", "30", "69.50%", "0.6549", "0.6974", "0.7658", "Conventional Baseline"],
        ["Logistic Reg. (Base)", "30", "61.75%", "0.5966", "0.6691", "0.7002", "Linear Baseline"],
        ["Genetic Alg. (GA) + RF", "4", "68.25%", "0.6432", "0.6867", "0.7514", "🌿 86.7% Pruned (Parsimonious)"],
        ["Particle Swarm (PSO) + RF", "9", "70.40%", "0.6602", "0.6974", "0.7630", "🏆 Champion Model (Best Acc)"]
    ]

    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
        p = cell.text_frame.paragraphs[0]
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_TITLE

    for r_idx, r in enumerate(rows):
        for c_idx, val in enumerate(r):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_MAIN
            p.font.size = Pt(9.5)
            if r_idx == 3: # PSO Champion
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(239, 246, 255)
                p.font.bold = True
                p.font.color.rgb = BLUE_PRIMARY
            elif r_idx == 2: # GA Parsimonious
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(240, 253, 244)
                p.font.bold = True
                p.font.color.rgb = EMERALD
            else:
                p.font.color.rgb = TEXT_BODY

    # Analytical Caption under Table
    tb_ac = s7.shapes.add_textbox(Inches(1.0), Inches(5.85), tbl_w - Inches(0.4), Inches(0.8))
    tf_ac = tb_ac.text_frame
    tf_ac.word_wrap = True
    p = tf_ac.paragraphs[0]
    p.text = "Key Takeaway: PSO + RF achieved the highest test accuracy (70.40%) and highest Macro F1 (0.6602) across all models, proving that nature-inspired feature pruning removes collinear noise and enhances generalization."
    p.font.name = FONT_MAIN
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    # Computational Profiling Card (Right, 35% width)
    prof_w = Inches(3.8)
    prof_x = Inches(8.7)
    create_card(s7, prof_x, Inches(2.8), prof_w, Inches(4.0), accent_color=INDIGO)

    tb_pr = s7.shapes.add_textbox(prof_x + Inches(0.2), Inches(2.95), prof_w - Inches(0.4), Inches(0.35))
    p = tb_pr.text_frame.paragraphs[0]
    p.text = "Computational Profiling & Latency"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    tb_pr_body = s7.shapes.add_textbox(prof_x + Inches(0.2), Inches(3.4), prof_w - Inches(0.4), Inches(3.2))
    tf_pr_body = tb_pr_body.text_frame
    tf_pr_body.word_wrap = True

    profs = [
        ("Optimization Search Runtime", "• BPSO (15 iter): 34.25s (2.1x faster)\n• GA (20 gen): 72.27s"),
        ("Model Training Latency", "• PSO Model: 0.402s (20% faster than baseline)\n• Baseline RF: 0.502s"),
        ("Per-Variant Inference Speed", "• PSO Model: 37.69 ms / variant\n• Baseline Model: 40.72 ms / variant"),
        ("Memory Footprint", "• Peak RAM < 280 MB during 65k training\n• Lightweight model artifacts (< 1.5 MB)")
    ]
    for idx, (label, val) in enumerate(profs):
        p = tf_pr_body.paragraphs[0] if idx == 0 else tf_pr_body.add_paragraph()
        p.text = label
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_TITLE
        p.space_before = Pt(4)

        p2 = tf_pr_body.add_paragraph()
        p2.text = val
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9)
        p2.font.color.rgb = TEXT_BODY

    # =========================================================================
    # SLIDE 8: Results & Comparison Charts
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_canvas(s8)
    add_top_nav(s8, "Visual Diagnostics", "Slide 8: Results & Visual Comparison Charts", "ROC curves and generational fitness convergence validating nature-inspired pruning", 8)

    c_w = Inches(5.7)
    c_h = Inches(4.3)

    # Plot 1: ROC Curves
    create_card(s8, Inches(0.8), Inches(1.5), c_w, c_h, accent_color=CORAL)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(1.65), c_w - Inches(0.4), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = "Receiver Operating Characteristic (ROC) Curves"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    roc_img = os.path.join("static", "plots", "roc_curves.png")
    if os.path.exists(roc_img):
        s8.shapes.add_picture(roc_img, Inches(1.0), Inches(2.1), width=c_w - Inches(0.4))

    # Plot 2: GA Convergence
    create_card(s8, Inches(6.8), Inches(1.5), c_w, c_h, accent_color=EMERALD)
    tb2 = s8.shapes.add_textbox(Inches(7.0), Inches(1.65), c_w - Inches(0.4), Inches(0.35))
    p = tb2.text_frame.paragraphs[0]
    p.text = "Genetic Algorithm Generational Convergence"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    ga_img = os.path.join("static", "plots", "ga_convergence.png")
    if os.path.exists(ga_img):
        s8.shapes.add_picture(ga_img, Inches(7.0), Inches(2.1), width=c_w - Inches(0.4))

    # Bottom Analytical Takeaway Card
    create_card(s8, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.9), accent_color=BLUE_PRIMARY, fill_color=BLUE_LIGHT)
    tb_bot = s8.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.333), Inches(0.7))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    p = tf_bot.paragraphs[0]
    p.text = "Biological & Algorithmic Validation Takeaways:"
    p.font.name = FONT_MAIN
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    p2 = tf_bot.add_paragraph()
    p2.text = "1. ROC curves demonstrate that GA (AUC = 0.751) and PSO (AUC = 0.763) match the full 30-feature baseline (AUC = 0.766) while discarding up to 86.7% of redundant genomic noise.\n2. GA convergence trajectory validates steady elitist fitness progression across 20 generations, reliably isolating the 4 core biomarkers: LoFtool, AF_TGP, Max_AF, and IMPACT_HIGH."
    p2.font.name = FONT_MAIN
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_BODY
    p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 9: Working Prototype & Live Clinical Decision Support
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_canvas(s9)
    add_top_nav(s9, "Prototype Demonstration", "Slide 9: Working Prototype: Clinical Decision Support", "Interactive web dashboard providing real-time discordance scoring and explainability", 9)

    # Left: Prototype Screenshot with simulated browser frame
    proto_w = Inches(6.8)
    proto_h = Inches(5.3)
    create_card(s9, Inches(0.8), Inches(1.5), proto_w, proto_h, accent_color=BLUE_PRIMARY)

    # Browser Mockup Bar
    bar = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.82), Inches(1.52), proto_w - Inches(0.04), Inches(0.38))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
    bar.line.fill.background()

    # Browser traffic dots
    dots = [
        (Inches(0.95), CORAL),
        (Inches(1.10), AMBER),
        (Inches(1.25), EMERALD)
    ]
    for dx, dc in dots:
        dot = s9.shapes.add_shape(MSO_SHAPE.OVAL, dx, Inches(1.65), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = dc
        dot.line.fill.background()

    # URL Pill inside browser
    url_pill = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.6), Inches(4.5), Inches(0.22))
    url_pill.fill.solid()
    url_pill.fill.fore_color.rgb = RGBColor(255, 255, 255)
    url_pill.line.color.rgb = BORDER_LIGHT
    url_pill.line.width = Pt(1)
    p = url_pill.text_frame.paragraphs[0]
    p.text = "http://127.0.0.1:5000/ — MutationHunter Clinical AI 2.0"
    p.font.name = FONT_MAIN
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    ss_img = os.path.join("screenshots", "04_clinical_predictor_live.png")
    if os.path.exists(ss_img):
        s9.shapes.add_picture(ss_img, Inches(0.95), Inches(2.0), width=proto_w - Inches(0.3))

    # Right: Prototype Features & Real-World Use
    feat_w = Inches(4.6)
    feat_x = Inches(7.9)
    create_card(s9, feat_x, Inches(1.5), feat_w, proto_h, accent_color=INDIGO)

    tb = s9.shapes.add_textbox(feat_x + Inches(0.2), Inches(1.65), feat_w - Inches(0.4), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = "System Capabilities & Applications"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_TITLE

    tb_feat = s9.shapes.add_textbox(feat_x + Inches(0.2), Inches(2.1), feat_w - Inches(0.4), Inches(4.5))
    tf_feat = tb_feat.text_frame
    tf_feat.word_wrap = True

    features = [
        ("Clinical Input Console", "Direct entry for CADD scores, LoFtool intolerance, allele frequencies, and consequence annotations."),
        ("1-Click Clinical Presets", "Instant evaluation chips for High-Conflict VUS, Consistent Benign, and Truncating Pathogenic variants."),
        ("Real-Time Conflict Alert", "Instant classification call with two-tone probability distribution bar and confidence rating."),
        ("Biological Explainability", "Generates human-readable clinical rationale detailing the biophysical drivers behind the conflict call."),
        ("Dual-Mode Architecture", "Flask Python API with browser-native JavaScript fallback engine for zero-failure deployment."),
        ("Real-World Clinical Impact", "Pre-screening triage in diagnostic sequencing pipelines, tumor boards, and VUS reinterpretation.")
    ]
    for idx, (title, desc) in enumerate(features):
        p = tf_feat.paragraphs[0] if idx == 0 else tf_feat.add_paragraph()
        p.text = f"•  {title}:"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_TITLE
        p.space_before = Pt(4)

        p2 = tf_feat.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9)
        p2.font.color.rgb = TEXT_BODY

    # =========================================================================
    # SLIDE 10: Conclusion & Future Scope
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_canvas(s10)
    add_top_nav(s10, "Summary & Future Roadmap", "Slide 10: Conclusion, Limitations & Future Scope", "Translational findings, engineering boundaries, and future research directions", 10)

    # 4 Quadrants
    create_card(s10, Inches(0.8), Inches(1.6), quad_w, quad_h, accent_color=EMERALD)
    tb = s10.shapes.add_textbox(Inches(1.0), Inches(1.8), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Key Findings & Contributions"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    findings = [
        "Nature-Inspired Computing resolves high-dimensional collinearity in genomics, outperforming full-feature baselines (70.40% Accuracy with PSO).",
        "GA successfully pruned 86.7% of attributes down to 4 validated biomarkers (LoFtool, AF_TGP, Max_AF, IMPACT_HIGH).",
        "Confirmed gene loss-of-function intolerance and population allele thresholds as primary determinants of clinical discordance."
    ]
    for f in findings:
        p = tf.add_paragraph()
        p.text = f"•  {f}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Quadrant 2: Project Limitations
    create_card(s10, Inches(6.8), Inches(1.6), quad_w, quad_h, accent_color=CORAL)
    tb = s10.shapes.add_textbox(Inches(7.0), Inches(1.8), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Project Limitations"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CORAL

    limits = [
        "Confined to single-nucleotide variants (SNVs) and small indels; complex structural translocations and copy number variants (CNVs) not modeled.",
        "Static cross-sectional ClinVar dataset; does not capture dynamic longitudinal assertion updates over time.",
        "Class imbalance (75% consistent vs 25% conflicting) requires continued tuning of cost-sensitive thresholds."
    ]
    for l in limits:
        p = tf.add_paragraph()
        p.text = f"•  {l}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Quadrant 3: Future Research Scope
    create_card(s10, Inches(0.8), Inches(4.35), quad_w, quad_h, accent_color=INDIGO)
    tb = s10.shapes.add_textbox(Inches(1.0), Inches(4.55), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Future Research & Engineering"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = INDIGO

    futures = [
        "Multi-Swarm Hybridization: Combining Ant Colony Optimization (ACO) with Artificial Immune Systems (AIS) for multi-objective search.",
        "Deep Protein Language Models: Benchmarking evolved subsets against AlphaMissense and ESM-1b embeddings.",
        "Clinical LIMS Integration: Developing HL7/FHIR compliant microservices for direct electronic health record workflow integration."
    ]
    for ft in futures:
        p = tf.add_paragraph()
        p.text = f"•  {ft}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Quadrant 4: Academic References
    create_card(s10, Inches(6.8), Inches(4.35), quad_w, quad_h, accent_color=BLUE_PRIMARY)
    tb = s10.shapes.add_textbox(Inches(7.0), Inches(4.55), quad_w - Inches(0.4), quad_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "4. Key Academic References"
    p.font.name = FONT_MAIN
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_PRIMARY

    refs = [
        ("Landrum et al. (2018)", "ClinVar: improving access to variant interpretations and supporting evidence.", "Nucleic Acids Res."),
        ("Rentzsch et al. (2019)", "CADD: predicting the deleteriousness of variants throughout the human genome.", "Nucleic Acids Res."),
        ("Kennedy & Eberhart (1995)", "Particle Swarm Optimization.", "Proc. IEEE Int. Conf. Neural Networks."),
        ("Goldberg, D. E. (1989)", "Genetic Algorithms in Search, Optimization and Machine Learning.", "Addison-Wesley.")
    ]
    for idx, (authors, title, journal) in enumerate(refs):
        p = tf.add_paragraph()
        p.text = f"[{idx+1}]  {authors}, \"{title}\" {journal}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(3)

    # Save presentation
    prs.save(output_path)
    print(f"Successfully generated high-design presentation: {output_path}")

if __name__ == "__main__":
    build_deck("MutationHunter_Presentation.pptx")
